import io

import docx
import pytest

from app.models.transcript import Segment
from app.pipeline.ingest import UnsupportedFile, parse
from app.pipeline.normalize import clean_text, normalize, render

VTT = b"""WEBVTT

1
00:00:01.000 --> 00:00:04.500
<v Priya>So, um, invoices land in the inbox.

2
00:00:04.500 --> 00:00:06.000
<v Priya>We enter them.

3
00:01:06.250 --> 00:01:09.000
Marco: Over five thousand comes to me.
"""

SRT = b"""1
00:00:01,500 --> 00:00:02,250
Bob: Hi there.

2
00:00:03,000 --> 00:00:04,000
No speaker here.
"""


def test_vtt_voice_tags_speaker_prefix_and_timestamps() -> None:
    segs = parse("call.VTT", VTT)
    assert [(s.speaker, s.start_ts, s.end_ts) for s in segs] == [
        ("Priya", 1.0, 4.5),
        ("Priya", 4.5, 6.0),
        ("Marco", 66.25, 69.0),
    ]
    assert segs[2].text == "Over five thousand comes to me."


def test_srt() -> None:
    segs = parse("call.srt", SRT)
    assert [(s.speaker, s.text, s.start_ts) for s in segs] == [
        ("Bob", "Hi there.", 1.5),
        (None, "No speaker here.", 3.0),
    ]


def test_txt_with_and_without_speakers() -> None:
    segs = parse("notes.txt", b"\xef\xbb\xbfDana: First we check.\n\nthen we ship it\nDr. Lee: Done.\n")
    assert [(s.speaker, s.text) for s in segs] == [
        ("Dana", "First we check."),
        (None, "then we ship it"),
        ("Dr. Lee", "Done."),
    ]
    assert segs[0].start_ts is None


def test_docx() -> None:
    d = docx.Document()
    d.add_paragraph("Dana: We create the record.")
    d.add_paragraph("")
    d.add_paragraph("Leo: IT orders the laptop.")
    buf = io.BytesIO()
    d.save(buf)
    assert [(s.speaker, s.text) for s in parse("t.docx", buf.getvalue())] == [
        ("Dana", "We create the record."),
        ("Leo", "IT orders the laptop."),
    ]


def test_unsupported_and_malformed() -> None:
    with pytest.raises(UnsupportedFile):
        parse("audio.mp3", b"")
    with pytest.raises(UnsupportedFile):
        parse("bad.vtt", b"not a vtt file")


@pytest.mark.parametrize(
    ("raw", "clean"),
    [
        ("A supplier invoice lands, and someone, um, enters it.", "A supplier invoice lands, and someone enters it."),
        ("Um, so we, uh, check it.", "So we check it."),
        ("You know, it is slow.", "It is slow."),
        ("It is, you know, broken.", "It is broken."),
        ("Do you know the code?", "Do you know the code?"),  # not a filler here
        ("Hmm.", ""),
        ("Uhh okay umbrella", "Okay umbrella"),  # 'umbrella' is not 'um'
    ],
)
def test_clean_text(raw: str, clean: str) -> None:
    assert clean_text(raw) == clean


def test_normalize_merges_same_speaker_and_assigns_ids() -> None:
    segs = normalize(
        parse("call.vtt", VTT) + [Segment(segment_id="x", speaker="Marco", start_ts=None, end_ts=None, text="Uh.")]
    )
    assert [(s.segment_id, s.speaker, s.start_ts, s.end_ts) for s in segs] == [
        ("s1", "Priya", 1.0, 6.0),
        ("s2", "Marco", 66.25, 69.0),  # filler-only fragment dropped, end_ts kept
    ]
    assert segs[0].text == "So invoices land in the inbox. We enter them."
    assert render(segs).splitlines()[0] == "[s1] Priya: So invoices land in the inbox. We enter them."


def test_unknown_speakers_are_never_merged() -> None:
    raw = [Segment(segment_id=str(i), speaker=None, start_ts=None, end_ts=None, text=t) for i, t in enumerate("ab")]
    assert [s.text for s in normalize(raw)] == ["A", "B"]


def test_unknown_speaker_label_is_never_merged() -> None:
    from app.models.transcript import UNKNOWN_SPEAKER

    raw = [
        Segment(segment_id=str(i), speaker=UNKNOWN_SPEAKER, start_ts=i, end_ts=i + 1, text=t)
        for i, t in enumerate("ab")
    ]
    assert [s.text for s in normalize(raw)] == ["A", "B"]
